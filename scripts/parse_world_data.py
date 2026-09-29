#!/usr/bin/env python3
"""Parse 'X-World' sheet of O राम - World Data.xlsx into Qumanity's
Global Timeline location JSON, using the same ID conventions as
SITAXML.xml (<Global_Continent id="O_AS"><Global_Country id="O_AS_IND">)
and global_core.py's ISO alpha-3 country codes.
"""
import json
import re
import sys

import openpyxl
import pycountry

SRC = "/Users/macmudgal/Desktop/Windows Data/sita foundation/SITA Project Data/O राम - World Data.xlsx"
OUT = "/Users/macmudgal/Desktop/Qumanity_DPG/static/data/global_countries.json"

# SITAXML / global_core.py continent codes (2-letter), matches the 7-continent
# marker set already plotted on the prototype globe.
CONTINENT_CODES = {
    "Asia": "AS",
    "Africa": "AF",
    "Europe": "EU",
    "North America": "NA",
    "South America": "SA",
    "Australia & Oceania": "OC",
    "Antarctica": "AN",
}

# Representative lat/lng per continent (same values already used in
# static/data/global.json) so the enhanced globe can keep plotting continent
# markers without guessing new coordinates.
CONTINENT_LATLNG = {
    "AS": (34.0, 100.6),
    "AF": (8.7, 34.5),
    "EU": (54.5, 15.2),
    "NA": (45.0, -100.0),
    "SA": (-15.0, -60.0),
    "OC": (-25.0, 135.0),
    "AN": (-82.0, 0.0),
}

# A handful of country-name spellings in the source sheet that don't match
# pycountry's official short/common names. Mapped by hand after the parse
# flagged them as unresolved.
NAME_FIXES = {
    "Ivory Coast": "Côte d'Ivoire",
    "Congo": "Congo, The Democratic Republic of the",
    "Republic of the Congo": "Congo",
    "Cape Verde": "Cabo Verde",
    "Swaziland": "Eswatini",
    "Burma": "Myanmar",
    "Laos": "Lao People's Democratic Republic",
    "North Korea": "Korea, Democratic People's Republic of",
    "South Korea": "Korea, Republic of",
    "Vatican City": "Holy See (Vatican City State)",
    "Micronesia": "Micronesia, Federated States of",
    "Vietnam": "Viet Nam",
    "Russia": "Russian Federation",
    "Syria": "Syrian Arab Republic",
    "Iran": "Iran, Islamic Republic of",
    "Brunei": "Brunei Darussalam",
    "Moldova": "Moldova, Republic of",
    "Bolivia": "Bolivia, Plurinational State of",
    "Venezuela": "Venezuela, Bolivarian Republic of",
    "Tanzania": "Tanzania, United Republic of",
    "United States": "United States of America",
    "South Sudan": "South Sudan",
    "The Bahamas": "Bahamas",
    "The Gambia": "Gambia",
    "East Timor": "Timor-Leste",
    "Palestine": "Palestine, State of",
    "Macedonia": "North Macedonia",
    "Czech Republic": "Czechia",
    "United Kingdom": "United Kingdom of Great Britain and Northern Ireland",
    "Nambia": "Namibia",
    "Zimbawe": "Zimbabwe",
    "Cote d'Ivoire (Ivory Coast)": "Côte d'Ivoire",
    "Turkey": "Türkiye",
    "Democratic Republic of Congo": "Congo, The Democratic Republic of the",
    "Republic of Congo": "Congo",
    "Bonaire": "Bonaire, Sint Eustatius and Saba",
    "Bonaire, Saba": "Bonaire, Sint Eustatius and Saba",
    "US Virgin Islands": "Virgin Islands, U.S.",
    "Pitcairn Islands": "Pitcairn",
    "Curacao": "Curaçao",
    "Sint Maarten": "Sint Maarten (Dutch part)",
}

# Not ISO 3166-1 sovereign-country codes but in real-world use / needed by
# the source sheet's own rows. Kept explicit rather than guessed.
MANUAL_CODES = {
    "Kosovo": "XKX",  # user-assigned, no official ISO 3166-1 code exists
}

# Rows in the source sheet that are NOT actually countries (data-quality
# issue in "O राम - World Data.xlsx", not something this parser should
# paper over). Flagged in the output instead of assigned a fake ISO3.
NOT_A_COUNTRY = {
    "Balochistan": "Province of Pakistan, not a separate country — appears as its own row in the source sheet.",
}


def resolve_iso3(name):
    name = name.strip()
    if name in MANUAL_CODES:
        return MANUAL_CODES[name]
    for candidate in (name, NAME_FIXES.get(name, "")):
        if not candidate:
            continue
        try:
            hit = pycountry.countries.lookup(candidate)
            return hit.alpha_3
        except LookupError:
            continue
    try:
        hits = pycountry.countries.search_fuzzy(NAME_FIXES.get(name, name))
        if hits:
            return hits[0].alpha_3
    except LookupError:
        pass
    return None


def slugify_fallback(name, used):
    base = re.sub(r"[^A-Za-z]", "", name).upper()[:3] or "XXX"
    code = base
    n = 1
    while code in used:
        code = (base[:2] + str(n))[:3]
        n += 1
    return code


def parse():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["X-World"]

    continents = []
    current = None
    unresolved = []
    used_codes_per_continent = {}

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        vals = [v for v in row if v is not None]
        if not vals:
            continue
        first = vals[0]
        if isinstance(first, str) and first.startswith("Y(X") and first != "Y(X)":
            name = vals[1]
            country_count = vals[2] if len(vals) > 2 else None
            city_count = vals[4] if len(vals) > 4 else None
            code = CONTINENT_CODES.get(name)
            if code is None:
                print("WARN unknown continent name:", name, file=sys.stderr)
                code = slugify_fallback(name, set())
            current = {
                "code": code,
                "id": "O_" + code,
                "name": name,
                "lat": CONTINENT_LATLNG.get(code, (0, 0))[0],
                "lng": CONTINENT_LATLNG.get(code, (0, 0))[1],
                "declared_country_count": country_count,
                "declared_city_count": city_count,
                "countries": [],
            }
            used_codes_per_continent[code] = {}
            continents.append(current)
        elif isinstance(first, str) and re.match(r"^0\(\d+-\d+\)$", first):
            if current is None:
                continue
            name = vals[2] if len(vals) > 2 else None
            city_count = vals[3] if len(vals) > 3 else 0
            if not name:
                continue
            iso3 = None if name in NOT_A_COUNTRY else resolve_iso3(name)
            if iso3 is None:
                iso3 = slugify_fallback(name, used_codes_per_continent[current["code"]])
                unresolved.append(name)

            seen = used_codes_per_continent[current["code"]]
            record_id = "O_" + current["code"] + "_" + iso3
            dup_note = None
            if iso3 in seen:
                seen[iso3] += 1
                dup_note = (
                    f"Source sheet has more than one row resolving to {iso3} "
                    f"in this continent (likely a duplicate/overlapping entry)."
                )
                record_id += "-" + str(seen[iso3])
            else:
                seen[iso3] = 1

            entry = {
                "code": iso3,
                "id": record_id,
                "name": name,
                "city_count": city_count,
            }
            if name in NOT_A_COUNTRY:
                entry["data_quality_flag"] = NOT_A_COUNTRY[name]
            elif dup_note:
                entry["data_quality_flag"] = dup_note
            current["countries"].append(entry)

    # Antarctica has no country rows in the source sheet; keep it as an
    # empty, non-selectable continent so the globe still shows all 7.
    if "AN" not in [c["code"] for c in continents]:
        continents.append(
            {
                "code": "AN",
                "id": "O_AN",
                "name": "Antarctica",
                "lat": CONTINENT_LATLNG["AN"][0],
                "lng": CONTINENT_LATLNG["AN"][1],
                "declared_country_count": 0,
                "declared_city_count": 0,
                "countries": [],
            }
        )

    total_countries = sum(len(c["countries"]) for c in continents)
    total_cities = sum(sum(co["city_count"] or 0 for co in c["countries"]) for c in continents)

    result = {
        "space": {"name": "Earth", "id": "O"},
        "source": "O राम - World Data.xlsx (X-World sheet)",
        "continent_count": len(continents),
        "country_count": total_countries,
        "city_count": total_cities,
        "continents": continents,
    }

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print("Continents:", len(continents))
    print("Countries total:", total_countries)
    print("Cities total:", total_cities)
    print("Unresolved ISO3 (fell back to generated code):", len(unresolved))
    for n in unresolved:
        print("  -", n)
    print("Wrote", OUT)


if __name__ == "__main__":
    parse()
